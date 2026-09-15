import os
import re

import fitz
import pytesseract

from pdf2image import convert_from_path

from PIL import Image

from docx import Document


# ============================================================
# 1. TESSERACT AND POPPLER PATHS
# ============================================================

pytesseract.pytesseract.tesseract_cmd = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)

POPPLER_PATH = (
    r"C:\Users\NADEEM ANSARI\Desktop\resume"
    r"\poppler-26.02.0\Library\bin"
)


# ============================================================
# 2. CLEAN TEXT
# ============================================================

def clean_text(text):

    if not text:

        return ""

    # Normalize line endings
    text = text.replace(
        "\r\n",
        "\n"
    )

    text = text.replace(
        "\r",
        "\n"
    )

    # Remove excessive spaces
    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    # Remove excessive blank lines
    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )

    # Remove spaces at beginning/end
    lines = []

    for line in text.split("\n"):

        line = line.strip()

        if line:

            lines.append(line)

    return "\n".join(lines)


# ============================================================
# 3. EXTRACT TEXT FROM NORMAL PDF
# ============================================================

def extract_text_from_normal_pdf(
    file_path
):

    document = fitz.open(
        file_path
    )

    all_text = []

    page_count = len(
        document
    )

    for page in document:

        page_text = page.get_text()

        if page_text:

            all_text.append(
                page_text
            )

    document.close()

    full_text = "\n\n".join(
        all_text
    )

    return (
        full_text,
        page_count
    )


# ============================================================
# 4. EXTRACT TEXT FROM SCANNED PDF USING OCR
# ============================================================

def extract_text_from_scanned_pdf(
    file_path
):

    print(
        "Converting PDF pages to images..."
    )

    images = convert_from_path(
        file_path,
        poppler_path=POPPLER_PATH
    )

    all_text = []

    page_count = len(
        images
    )

    for page_number, image in enumerate(
        images
    ):

        print(
            f"Processing page "
            f"{page_number + 1} "
            f"of {page_count}"
        )

        page_text = pytesseract.image_to_string(
            image,
            config="--psm 4"
        )

        all_text.append(
            page_text
        )

    full_text = "\n\n".join(
        all_text
    )

    return (
        full_text,
        page_count
    )


# ============================================================
# 5. EXTRACT TEXT FROM PDF
# ============================================================

def extract_text_from_pdf(
    file_path
):

    # First try normal PDF text extraction

    normal_text, page_count = (
        extract_text_from_normal_pdf(
            file_path
        )
    )

    cleaned_normal_text = clean_text(
        normal_text
    )

    # If enough text was extracted,
    # consider it a normal text PDF

    if len(cleaned_normal_text) > 50:

        return {
            "text": cleaned_normal_text,
            "page_count": page_count,
            "method": "PDF Text Extraction"
        }

    # Otherwise use OCR

    print(
        "Very little text detected."
    )

    print(
        "Switching to OCR..."
    )

    ocr_text, page_count = (
        extract_text_from_scanned_pdf(
            file_path
        )
    )

    cleaned_ocr_text = clean_text(
        ocr_text
    )

    return {
        "text": cleaned_ocr_text,
        "page_count": page_count,
        "method": "PDF OCR"
    }


# ============================================================
# 6. EXTRACT TEXT FROM DOCX
# ============================================================

def extract_text_from_docx(
    file_path
):

    document = Document(
        file_path
    )

    all_text = []

    # Extract paragraphs

    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            all_text.append(
                paragraph.text
            )

    # Extract tables

    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                if cell.text.strip():

                    row_text.append(
                        cell.text.strip()
                    )

            if row_text:

                all_text.append(
                    " | ".join(
                        row_text
                    )
                )

    full_text = "\n".join(
        all_text
    )

    cleaned_text = clean_text(
        full_text
    )

    return {
        "text": cleaned_text,
        "page_count": 1,
        "method": "DOCX Text Extraction"
    }


# ============================================================
# 7. EXTRACT TEXT FROM TXT
# ============================================================

def extract_text_from_txt(
    file_path
):

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:

        text = file.read()

    cleaned_text = clean_text(
        text
    )

    return {
        "text": cleaned_text,
        "page_count": 1,
        "method": "TXT Text Extraction"
    }


# ============================================================
# 8. EXTRACT TEXT FROM IMAGE
# ============================================================

def extract_text_from_image(
    file_path
):

    image = Image.open(
        file_path
    )

    text = pytesseract.image_to_string(
        image,
        config="--psm 4"
    )

    cleaned_text = clean_text(
        text
    )

    return {
        "text": cleaned_text,
        "page_count": 1,
        "method": "Image OCR"
    }


# ============================================================
# 9. MAIN FILE PROCESSOR
# ============================================================

def extract_text_from_file(
    file_path,
    file_extension
):

    file_extension = (
        file_extension.lower()
    )

    if file_extension == "pdf":

        return extract_text_from_pdf(
            file_path
        )

    elif file_extension == "docx":

        return extract_text_from_docx(
            file_path
        )

    elif file_extension == "txt":

        return extract_text_from_txt(
            file_path
        )

    elif file_extension in [
        "jpg",
        "jpeg",
        "png"
    ]:

        return extract_text_from_image(
            file_path
        )

    else:

        raise ValueError(
            "Unsupported file type"
        )
